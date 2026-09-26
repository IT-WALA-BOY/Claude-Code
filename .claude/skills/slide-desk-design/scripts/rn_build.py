#!/usr/bin/env python3
"""RecNation Slides - deterministic PPTX builder.

    python rn_build.py deck.json out.pptx

The deck JSON holds content only (verbatim from the source). Every visual
decision lives here, once, so every slide from every teammate comes out the
same. Units: the design system is authored in CSS px on a 1280 x 720 canvas.
1 px = 1/96 in = 9525 EMU, and font px * 0.75 = pt. Nothing is eyeballed.
"""
import json, os, re, sys
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "..", "assets")
TOK = json.load(open(os.path.join(HERE, "tokens.json")))
C = TOK["color"]

W, H = 1280, 720
PAD_T, PAD_S, PAD_B = 40, 52, 16
CW = W - 2 * PAD_S  # content width 1176

# PowerPoint cannot select weight 800 from a "bold" flag. Each weight is its
# own face name; this table is why v1 decks came out in the wrong weights.
FACE = {500: ("Plus Jakarta Sans Medium", False), 600: ("Plus Jakarta Sans SemiBold", False),
        700: ("Plus Jakarta Sans", True), 800: ("Plus Jakarta Sans ExtraBold", False)}
TTF = {500: "Medium", 600: "SemiBold", 700: "Bold", 800: "ExtraBold"}
_fc = {}
SAFE = 3  # px per column: renderers break lines a hair earlier than the font file says


def font(weight, px):
    k = (weight, round(px * 4))
    if k not in _fc:
        _fc[k] = ImageFont.truetype(os.path.join(ASSETS, "fonts", f"PlusJakartaSans-{TTF[weight]}.ttf"), px * 4)
    return _fc[k]


def tw(text, weight, px):
    """Rendered width in px, measured with the real font file."""
    return font(weight, px).getlength(str(text)) / 4


def wrap(text, weight, px, width):
    lines, cur = [], ""
    for word in str(text).split(" "):
        t = (cur + " " + word).strip()
        if cur and tw(t, weight, px) > width:
            lines.append(cur); cur = word
        else:
            cur = t
    lines.append(cur)
    return lines


def E(px):
    return Emu(int(round(px * 9525)))


def rgb(hexs):
    return RGBColor.from_string(hexs.lstrip("#").upper())


# Density ladder (tables). body px, line px, header px, cell side padding px
DENSITY = [("roomy", 14.5, 38, 11.5, 18), ("open", 13.5, 30, 11, 16), ("standard", 12.5, 23, 10.5, 10),
           ("compact", 11.5, 20, 10, 14), ("tight", 10.4, 18, 9.4, 13), ("micro", 9.4, 15, 8.6, 11)]
CHIP_PX = {"roomy": 12, "open": 11, "standard": 11, "compact": 9.6, "tight": 9.2, "micro": 8.6}


# ------------------------------------------------------------------ primitives
def text(slide, x, y, w, h, runs, px, weight=500, color="ink", align="l", anchor="t", lh=None, ls=None):
    """runs: str or list of (text, weight, colorkey). One paragraph per '\n'."""
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(runs, str):
        runs = [(runs, weight, color)]
    paras = [[]]
    for t, wt, col in runs:
        parts = t.split("\n")
        for i, part in enumerate(parts):
            if i:
                paras.append([])
            if part:
                paras[-1].append((part, wt, col))
    for i, pr in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if lh:
            p.line_spacing = Pt(lh * 0.75)
        for t, wt, col in pr:
            r = p.add_run()
            r.text = t
            style_run(r, wt, px, col, ls)
        end_size(p, px)
    return tb


def style_run(r, weight, px, color, ls=None):
    face, bold = FACE[weight]
    f = r.font
    f.name = face
    f.size = Pt(px * 0.75)
    f.bold = bold
    f.color.rgb = rgb(C.get(color, color))
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", face)
    if ls:
        rPr.set("spc", str(int(ls * px * 75)))  # em -> 1/100 pt


def end_size(para, px):
    """Empty cells otherwise fall back to 18pt and silently inflate every row."""
    p = para._p
    e = p.find(qn("a:endParaRPr"))
    if e is None:
        e = etree.SubElement(p, qn("a:endParaRPr"))
    e.set("sz", str(int(px * 75)))
    e.set("lang", "en-US")


def box(slide, x, y, w, h, fill=None, line=None, radius=None, shadow=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    if radius:
        shp.adjustments[0] = radius / min(w, h)
    if fill:
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb(C.get(fill, fill))
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(C.get(line, line)); shp.line.width = E(1)
    else:
        shp.line.fill.background()
    spPr = shp._element.spPr
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    if shadow:  # elevation-card: 0 3px 10px rgba(0,26,74,.04)
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(9525 * 10), dist=str(9525 * 3), dir="5400000", algn="t", rotWithShape="0")
        c = etree.SubElement(sh, qn("a:srgbClr"), val="001A4A")
        etree.SubElement(c, qn("a:alpha"), val="4000")
    shp.text_frame.text = ""
    # kill theme style so nothing inherits the Office blue
    st = shp._element.find(qn("p:style"))
    if st is not None:
        shp._element.remove(st)
    return shp


# ------------------------------------------------------------------ tables
def set_border(cell, side, color, width_px):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {"L": "a:lnL", "R": "a:lnR", "T": "a:lnT", "B": "a:lnB"}[side]
    old = tcPr.find(qn(tag))
    if old is not None:
        tcPr.remove(old)
    ln = etree.Element(qn(tag), w=str(int(width_px * 9525)) if color else "0", cap="flat", cmpd="sng", algn="ctr")
    if color:
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=C.get(color, color).lstrip("#"))
    else:
        etree.SubElement(ln, qn("a:noFill"))
    # schema order: lnL lnR lnT lnB ... fill
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
    idx = order.index(tag)
    pos = 0
    for i, ch in enumerate(list(tcPr)):
        if ch.tag in [qn(t) for t in order[:idx]]:
            pos = i + 1
    tcPr.insert(pos, ln)


def fill_cell(cell, color):
    if color:
        cell.fill.solid(); cell.fill.fore_color.rgb = rgb(C.get(color, color))
    else:
        cell.fill.background()


def is_neg(s):
    s = str(s).strip()
    return s.startswith("(") or s.startswith("−") or s.startswith("-")


def chip_text(s):
    """Chips take the minus sign, plain cells keep parentheses."""
    s = str(s).strip()
    m = re.match(r"^\((.*)\)(.*)$", s)
    return "−" + m.group(1) + m.group(2) if m else s.replace("-", "−", 1) if s.startswith("-") else s


def chip_kind(s):
    s = str(s).strip()
    if s in ("", "—", "–"):
        return None
    if re.fullmatch(r"[+−-]?0(\.0+)?\s*(%|pp|bps)?", s.replace("(", "").replace(")", "")):
        return "zero"
    return "dn" if is_neg(s) else "up"


class Table:
    """One engine for every table type. kind: statement | breakdown | series | experiment.

    spec: {kind, heading?, subheading?, note?, corner, groups:[{label, cols:[label..]}],
           rows:[{label, cells:[...], style: member|total|pct|plain|control|win|now}],
           chips:[flat col idx], ramp:{col, dir:'up'|'down'}, fill_width: bool}
    """

    def __init__(self, spec):
        self.s = spec
        self.kind = spec.get("kind", "statement")
        self.groups = spec["groups"]
        self.ncols = sum(len(g["cols"]) for g in self.groups)
        self.has_groups = any(g.get("label") for g in self.groups) and len(self.groups) > 1
        self.gutter = self.has_groups and (len(self.groups) >= 3 or self.ncols > 8)
        self.chips = set(spec.get("chips", []))
        self.ramp = spec.get("ramp")
        if self.ramp:
            vals = [self._num(r["cells"][self.ramp["col"]]) for r in spec["rows"] if r.get("style") != "total"]
            vals = [v for v in vals if v is not None]
            self.rlo, self.rhi = min(vals), max(vals)

    @staticmethod
    def _num(s):
        t = re.sub(r"[^0-9.\-−()]", "", str(s)).replace("−", "-")
        neg = t.startswith("(")
        t = t.strip("()")
        try:
            v = float(t)
        except ValueError:
            return None
        return -v if neg else v

    def band(self, s):
        v = self._num(s)
        if v is None or self.rhi == self.rlo:
            return 3
        f = (v - self.rlo) / (self.rhi - self.rlo)
        if self.ramp.get("dir", "up") == "down":
            f = 1 - f
        return min(5, int(f * 5) + 1)

    def row_weight(self, st):
        return {"total": 800, "now": 800, "win": 700, "control": 600, "pct": 500}.get(st, 500)

    def heading_h(self):
        h = 0
        if self.s.get("heading"):
            h += 17
        if self.s.get("subheading"):
            h += 15
        return h + (6 if h else 0)

    def note_h(self, w):
        h = 0
        if self.s.get("note"):
            h += 6 + 13 * len(wrap(self.s["note"], 500, 9.6, w))
        if self.kind == "experiment" and self.s.get("caveat"):
            h += 8 + 12 + 13 * len(wrap(self.s["caveat"], 600, 10, w - 20))
        return h

    def col_widths(self, d):
        name, body, line, hdr, pad = d
        rows = self.s["rows"]
        lw = max([tw(r["label"], self.row_weight(r.get("style")) if r.get("style") in ("total", "now", "win") else 600, body) +
                  (30 if self.kind == "statement" and r.get("style", "member") == "member" else 0) for r in rows] + [tw(self.s.get("corner", ""), 700, hdr)])
        widths = [lw + 2 * pad + 12]
        flat = 0
        for gi, g in enumerate(self.groups):
            gcols = []
            for c in g["cols"]:
                cw = tw(c, 700, hdr)
                for r in rows:
                    v = r["cells"][flat] if flat < len(r["cells"]) else ""
                    wt = 800 if r.get("style") in ("total", "now") else 500
                    if flat in self.chips:
                        cw = max(cw, tw(chip_text(v), 700, CHIP_PX[name]) + 14)
                    else:
                        cw = max(cw, tw(v, wt, body))
                gcols.append(cw + 2 * pad + SAFE)
                flat += 1
            if g.get("label"):  # group label must fit over its columns
                need = tw(g["label"], 800, 10.5) + 16
                if sum(gcols) < need:
                    gcols = [c * need / sum(gcols) for c in gcols]
            widths.append(gcols)
        return widths

    def measure(self, d, avail_w):
        cw = self.col_widths(d)
        natural = cw[0] + sum(sum(g) for g in cw[1:]) + (10 * (len(self.groups) - 1) if self.gutter else 0)
        name, body, line, hdr, pad = d
        nh = (line if name in ("roomy", "open") else max(line, 20)) if False else line
        h = self.heading_h() + (26 if self.has_groups else 0) + max(24, hdr * 2.2) + len(self.s["rows"]) * line + 2
        w = avail_w if self.s.get("fill_width") else natural
        return natural, w, h + self.note_h(w)

    def draw(self, slide, x, y, w, d):
        s, name = self.s, d[0]
        body, line, hdr, pad = d[1], d[2], d[3], d[4]
        cw = self.col_widths(d)
        natural = cw[0] + sum(sum(g) for g in cw[1:]) + (10 * (len(self.groups) - 1) if self.gutter else 0)
        extra = (w - natural) / max(1, self.ncols + 1)  # breathe evenly
        yy = y
        if s.get("heading"):
            text(slide, x, yy, w, 17, s["heading"], 11.5, 800, "brand-navy"); yy += 17
        if s.get("subheading"):
            text(slide, x, yy, w, 15, s["subheading"], 10.5, 500, "ink-subhead"); yy += 15
        if s.get("heading") or s.get("subheading"):
            yy += 6
        # column plan: ('label'|'data'|'sp', width, group index, flat index)
        plan = [("label", cw[0] + extra, -1, -1)]
        flat = 0
        for gi, g in enumerate(cw[1:]):
            if gi and self.gutter:
                plan.append(("sp", 10, -1, -1))
            for c in g:
                plan.append(("data", c + extra, gi, flat)); flat += 1
        hdr_rows = 2 if self.has_groups else 1
        nrows = hdr_rows + len(s["rows"])
        hrow = max(24, hdr * 2.2)
        th = (26 if self.has_groups else 0) + hrow + len(s["rows"]) * line
        tw_total = sum(p[1] for p in plan)
        # container: rounded, bordered, elevated; table sits inside
        box(slide, x, yy, tw_total, th + 2, fill="surface", line="line", radius=8, shadow=True)
        gf = slide.shapes.add_table(nrows, len(plan), E(x + 1), E(yy + 1), E(tw_total - 2), E(th))
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        for a in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
            tblPr.set(a, "0")
        sid = tblPr.find(qn("a:tableStyleId"))
        if sid is not None:
            sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid" base
        for i, p in enumerate(plan):
            tbl.columns[i].width = E(p[1] - (2 / len(plan)))
        gf.width = E(tw_total - 2)
        if self.has_groups:
            tbl.rows[0].height = E(26)
        tbl.rows[hdr_rows - 1].height = E(hrow)
        washes = ["wash-1", "wash-2", "wash-3"]
        ramp_bg = {i: (C[f"ramp-{i}-bg"], C[f"ramp-{i}-text"]) for i in range(1, 6)}

        def put(cell, t, px, wt, col, align, lpad=None):
            cell.margin_left = E(pad if lpad is None else lpad); cell.margin_right = E(pad)
            cell.margin_top = cell.margin_bottom = 0
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            para = tf.paragraphs[0]
            para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            # Plus Jakarta Sans carries tall win metrics (1.65 em). Left on "single"
            # spacing PowerPoint grows every row ~30% and the table runs off the slide.
            para.line_spacing = Pt(px * 0.75 * 1.2)
            # an empty run is dropped by some renderers, which then size the row
            # from the 18pt default; a sized space keeps every row at its step
            r = para.add_run(); r.text = str(t) if str(t) else " "
            style_run(r, wt, px, col)
            end_size(para, px)

        def grid(cell, ri, ci, last_r, last_c, header=False):
            set_border(cell, "L", None, 0); set_border(cell, "T", None, 0)
            set_border(cell, "R", None if last_c else "line-grid", 1)
            set_border(cell, "B", None if last_r else "line-grid", 1)

        # header row 1: group labels (white, brand 800, centred, merged)
        ci = 0
        if self.has_groups:
            for pi, p in enumerate(plan):
                c = tbl.cell(0, pi); fill_cell(c, "surface")
                if p[0] == "sp":
                    put(c, "", 10.5, 800, "brand", "l", lpad=0); c.margin_right = 0
                for side in "LRTB":
                    set_border(c, side, None, 0)
            pi = 1
            for gi, g in enumerate(self.groups):
                if gi and self.gutter:
                    pi += 1
                n = len(g["cols"])
                c = tbl.cell(0, pi)
                if n > 1:
                    c.merge(tbl.cell(0, pi + n - 1))
                put(c, g.get("label", ""), 10.5, 800, "brand", "c")
                for k in range(n):
                    set_border(tbl.cell(0, pi + k), "B", "line-grid", 1)
                pi += n
            put(tbl.cell(0, 0), "", 10.5, 800, "brand", "l")
        # header row 2: column labels on one continuous grey band
        hr = hdr_rows - 1
        for pi, p in enumerate(plan):
            c = tbl.cell(hr, pi)
            fill_cell(c, "surface-header")
            if p[0] == "label":
                put(c, s.get("corner", ""), hdr, 700, "ink-muted", "l")
            elif p[0] == "data":
                gcols = self.groups[p[2]]["cols"]
                idx = p[3] - sum(len(g["cols"]) for g in self.groups[:p[2]])
                put(c, gcols[idx], hdr, 700, "ink-muted", "r" if self.kind != "reference" else "l")
            else:
                put(c, "", hdr, 700, "ink-muted", "l", lpad=0); c.margin_right = 0
            for side in "LRT":
                set_border(c, side, None, 0)
            set_border(c, "B", "line-grid", 1)
        # body
        rows = s["rows"]
        for ri, r in enumerate(rows):
            R = hdr_rows + ri
            tbl.rows[R].height = E(line)
            st = r.get("style", "member" if self.kind == "statement" else "plain")
            wt = self.row_weight(st)
            last_r = ri == len(rows) - 1
            rowfill = {"total": "brand-tint-row", "now": "brand-tint-row", "win": "brand-tint-row"}.get(st)
            for pi, p in enumerate(plan):
                c = tbl.cell(R, pi)
                last_c = pi == len(plan) - 1
                if p[0] == "sp":
                    fill_cell(c, rowfill or "surface")
                    c.margin_left = c.margin_right = 0
                    for side in "LRTB":
                        set_border(c, side, None, 0)
                    set_border(c, "B", None if last_r else "line-grid", 1)
                    put(c, "", body, 500, "ink", "l", lpad=0); c.margin_right = 0
                    continue
                if p[0] == "label":
                    lab_fill = {"total": "brand-tint-row", "now": "brand-tint-label", "win": "brand-tint-label"}.get(st, "surface-label-col")
                    fill_cell(c, lab_fill)
                    col = {"member": "ink-secondary", "pct": "ink-tertiary", "control": "ink-subhead"}.get(st, "ink")
                    lw = {"member": 600, "plain": 600, "pct": 500}.get(st, wt)
                    ind = pad + (30 if st == "member" and self.kind == "statement" else 0)
                    put(c, r["label"], 11 if st == "pct" else body, lw, col, "l", lpad=ind)
                    if st == "win" and s.get("best_pill", True):
                        pass
                else:
                    v = r["cells"][p[3]] if p[3] < len(r["cells"]) else ""
                    f = rowfill
                    if not f and self.gutter:
                        f = washes[p[2] % 3]
                    if st == "total" and self.gutter:
                        f = washes[p[2] % 3] + "-total"
                    col = {"pct": "ink-tertiary", "control": "ink-subhead"}.get(st, "ink")
                    px = 11 if st == "pct" else body
                    cw_ = wt
                    if p[3] in self.chips and chip_kind(v):
                        k = chip_kind(v)  # a cell treatment owns its colour
                        f = {"up": "chip-up-bg", "dn": "chip-down-bg", "zero": "chip-zero-bg"}[k]
                        col = {"up": "positive", "dn": "chip-down-text", "zero": "ink-muted"}[k]
                        v = chip_text(v); px = CHIP_PX[name]; cw_ = max(700, wt)
                    elif self.ramp and p[3] == self.ramp["col"] and st != "total":
                        b = self.band(v)
                        f, col = ramp_bg[b]; cw_ = 700
                    elif is_neg(v) and st not in ("pct",) and s.get("red_negatives", True):
                        col = "negative"
                    fill_cell(c, f or "surface")
                    put(c, v, px, cw_, col, "r")
                grid(c, R, pi, last_r, last_c)
                if st == "control" and ri > 0:
                    set_border(c, "T", "line-control", 2)
                    set_border(tbl.cell(R - 1, pi), "B", "line-control", 2)
        yy += th + 2
        if s.get("note"):
            yy += 6
            text(slide, x, yy, w, 13 * len(wrap(s["note"], 500, 9.6, w)), [(s["note"], 500, "ink-faint")], 9.6, lh=13)
            yy += 13 * len(wrap(s["note"], 500, 9.6, w))
        if self.kind == "experiment" and s.get("caveat"):
            yy += 8
            lines = wrap(s["caveat"], 600, 10, w - 20)
            hh = 12 + 13 * len(lines)
            box(slide, x, yy, w, hh, fill="caveat-bg", radius=6)
            box(slide, x, yy, 3, hh, fill="caveat-bar")
            text(slide, x + 12, yy + 6, w - 20, hh - 12, [(s["caveat"], 600, "caveat-text")], 10, lh=13)


# ------------------------------------------------------------------ blocks
class Cards:
    def __init__(self, spec):
        self.s = spec

    def measure(self, d, avail_w):
        fig = 30 if d[0] == "micro" else 38
        return avail_w, avail_w, 16 + 14 + 4 + fig + 6 + 16 + 18

    def draw(self, slide, x, y, w, d):
        items = self.s["items"]
        n = len(items); gap = 16
        cw = (w - gap * (n - 1)) / n
        fig, fl = (30, 35) if d[0] == "micro" else (38, 44)
        h = self.measure(d, w)[2]
        for i, it in enumerate(items):
            cx = x + i * (cw + gap)
            box(slide, cx, y, cw, h, fill="surface", line="line-card", radius=8, shadow=True)
            text(slide, cx + 20, y + 16, cw - 40, 14, it["label"].upper(), 10.5, 800, "ink-label", ls=0.08)
            text(slide, cx + 20, y + 16 + 18, cw - 40, fl, it["value"], fig, 800, "brand" if it.get("accent") else "ink", ls=-0.03)
            ctx = it.get("context", "")
            text(slide, cx + 20, y + 16 + 18 + fl + 4, cw - 40, 16, ctx, 11.5, 500, "ink-muted")


class Bullets:
    def __init__(self, spec):
        self.s = spec

    def layout(self, w, load):
        cols = 1 if load == "open" else 2
        px, lh = (13, 19) if cols == 1 else (12, 17)
        colw = (w - 34 * (cols - 1)) / cols
        items = self.s["items"]
        per = -(-len(items) // cols)
        colsets = [items[i * per:(i + 1) * per] for i in range(cols)]
        heights = [sum(len(wrap(t, 500, px, colw - 16)) * lh + 3 for t in cs) for cs in colsets]
        return cols, px, lh, colw, colsets, max(heights)

    def draw(self, slide, x, y, w, load):
        cols, px, lh, colw, colsets, _ = self.layout(w, load)
        for ci, cs in enumerate(colsets):
            yy = y
            for t in cs:
                n = len(wrap(t, 500, px, colw - 16))
                dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(x + ci * (colw + 34) + 2), E(yy + lh / 2 - 2), E(4), E(4))
                dot.fill.solid(); dot.fill.fore_color.rgb = rgb(C["brand"]); dot.line.fill.background()
                text(slide, x + ci * (colw + 34) + 16, yy, colw - 16, n * lh, [(t, 500, "ink-secondary")], px, lh=lh)
                yy += n * lh + 3


class Row:
    """Blocks side by side, pair-gap apart. Width shared in proportion to natural width."""
    GAP = 22

    def __init__(self, spec):
        self.s = spec
        self.items = [BLOCKS[b["type"]](b) for b in spec["items"]]

    def _split(self, d, avail_w):
        nats = [b.measure(d, avail_w)[0] for b in self.items]
        room = avail_w - self.GAP * (len(self.items) - 1)
        return nats, [room * n / sum(nats) for n in nats]

    def measure(self, d, avail_w):
        nats, ws = self._split(d, avail_w)
        nat = sum(nats) + self.GAP * (len(self.items) - 1)
        return nat, avail_w, max(b.measure(d, w)[2] for b, w in zip(self.items, ws))

    def draw(self, slide, x, y, w, d):
        nats, ws = self._split(d, w)
        for b, bw in zip(self.items, ws):
            b.s["fill_width"] = True
            b.draw(slide, x, y, bw, d)
            x += bw + self.GAP


BLOCKS = {"table": Table, "cards": Cards, "row": Row}


def takeaway_h(tk, rec, w, load):
    px, lh, py = {"open": (16.5, 23, 13), "standard": (15, 22, 12), "dense": (15, 22, 11)}[load]
    if rec:
        half = (w - 14) / 2
        a = len(wrap(tk, 700, px, half - 40)) * lh
        rpx, rlh = (14, 19) if load == "dense" else (px, lh)
        b = len(wrap(rec, 700, rpx, half - 44)) * rlh
        return max(a, b) + 2 * py
    return len(wrap(tk, 700, px, w - 40)) * lh + 2 * py


def draw_takeaway(slide, x, y, w, h, tk, rec, load):
    px, lh, py = {"open": (16.5, 23, 13), "standard": (15, 22, 12), "dense": (15, 22, 11)}[load]
    tw_ = (w - 14) / 2 if rec else w
    box(slide, x, y, tw_, h, fill="brand", radius=8)
    runs = []
    m = re.match(r"^(KEY TAKEAWAY)\s*(.*)$", tk, re.S)
    if m:
        runs = [(m.group(1) + "   ", 800, "brand-label-on-solid"), (m.group(2), 700, "surface")]
    else:
        runs = [(tk, 700, "surface")]
    text(slide, x + 20, y + py, tw_ - 40, h - 2 * py, runs, px, lh=lh, ls=-0.02, anchor="m")
    if rec:
        rx = x + tw_ + 14
        rpx, rlh = (14, 19) if load == "dense" else (px, lh)
        box(slide, rx, y, tw_, h, fill="brand-pale", radius=8)
        box(slide, rx, y, 4, h, fill="brand")
        m = re.match(r"^(RECOMMENDATION)\s*(.*)$", rec, re.S)
        runs = [(m.group(1) + "   ", 800, "brand"), (m.group(2), 700, "brand-deep")] if m else [(rec, 700, "brand-deep")]
        text(slide, rx + 24, y + py, tw_ - 44, h - 2 * py, runs, rpx, lh=rlh, ls=-0.02, anchor="m")


# ------------------------------------------------------------------ slides
def header(slide, s):
    if s.get("eyebrow"):
        text(slide, PAD_S, PAD_T, 800, 16, s["eyebrow"], 12, 700, "brand")
    text(slide, PAD_S, PAD_T + 18, CW - 130, 38, s["title"], 30, 800, "ink", ls=-0.02)
    y = PAD_T + 18 + 38
    if s.get("subtitle"):
        text(slide, PAD_S, y + 2, CW - 130, 19, s["subtitle"], 13, 500, "ink-secondary", ls=-0.02)
        y += 21
    slide.shapes.add_picture(os.path.join(ASSETS, "logo-colour.png"), E(W - PAD_S - 104 + 11), E(PAD_T + 11), E(82))
    return y


def footer(slide, s, n, deck):
    y = H - PAD_B - 14
    text(slide, PAD_S, y, 300, 14, deck.get("copyright", "© RecNation Storage"), 10, 500, "ink-faint")
    src = s.get("source", "")
    if src:
        text(slide, PAD_S + 300, y, CW - 340, 14, src, 10, 500, "ink-faint", align="r")
    text(slide, W - PAD_S - 30, y, 30, 14, str(n), 9.5, 700, "ink-muted", align="r")


def build_content(prs, s, n, deck, report):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    hy = header(slide, s)
    top = hy + 18
    bottom = H - PAD_B - 14 - 14  # footer top minus takeaway-air
    blocks = [BLOCKS[b["type"]](b) for b in s.get("blocks", [])]
    rows = sum(len(b.s.get("rows", [])) for b in blocks if isinstance(b, Table))
    maxcols = max([b.ncols for b in blocks if isinstance(b, Table)] + [0])
    load = s.get("load") or ("dense" if rows >= 25 or maxcols >= 14 else "standard" if rows >= 10 else "open")
    tk, rec = s.get("takeaway"), s.get("recommendation")
    tkh = takeaway_h(tk, rec, CW, load) if tk else 0
    ty = bottom - tkh
    bl = Bullets({"items": s["bullets"]}) if s.get("bullets") else None
    blh = bl.layout(CW, load)[5] if bl else 0
    zone_bottom = ty - (9 if tk else 0) - (blh + 9 if bl else 0)
    zone = zone_bottom - top
    chosen = None
    for d in DENSITY:
        hs, ok = [], True
        for b in blocks:
            nat, w, h = b.measure(d, CW)
            if nat > CW + 0.5:
                ok = False
            hs.append(h)
        total = sum(hs) + sum((19 if isinstance(blocks[i], Cards) else 9) for i in range(len(blocks) - 1))
        if ok and total <= zone * 0.94:
            chosen = (d, hs, total); break
    if not chosen:
        d = DENSITY[-1]
        hs = [b.measure(d, CW)[2] for b in blocks]
        chosen = (d, hs, sum(hs))
        report.append(f"slide {n}: DOES NOT FIT even at Micro, split it at a parent boundary")
    d, hs, total = chosen
    report.append(f"slide {n}: load={load} density={d[0]} content {total:.0f}px of {zone:.0f}px")
    y = top + max(0, (zone - total) / 2)  # centred, never pooling underneath
    for i, b in enumerate(blocks):
        nat, w, h = b.measure(d, CW)
        b.draw(slide, PAD_S, y, w, d)
        y += h + (19 if isinstance(b, Cards) else 9)
    if bl:
        bl.draw(slide, PAD_S, zone_bottom + 9 - (9 if not tk else 0) + 0, CW, load)
    if tk:
        draw_takeaway(slide, PAD_S, ty, CW, tkh, tk, rec, load)
    footer(slide, s, n, deck)
    if s.get("notes"):
        slide.notes_slide.notes_text_frame.text = s["notes"]


def build_cover(prs, s):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box(slide, 0, 0, W, H, fill="ink")
    slide.shapes.add_picture(os.path.join(ASSETS, "cover-hero.jpg"), 0, 0, E(W), E(H))
    tint = box(slide, 0, 0, W, H, fill="ink")
    # ink tint over the photo, strongest at the foot where the title sits
    sf = tint._element.spPr.find(qn("a:solidFill"))
    grad = etree.fromstring(
        '<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" rotWithShape="1"><a:gsLst>'
        '<a:gs pos="0"><a:srgbClr val="030A17"><a:alpha val="35000"/></a:srgbClr></a:gs>'
        '<a:gs pos="45000"><a:srgbClr val="030A17"><a:alpha val="10000"/></a:srgbClr></a:gs>'
        '<a:gs pos="100000"><a:srgbClr val="030A17"><a:alpha val="80000"/></a:srgbClr></a:gs>'
        '</a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>')
    sf.addprevious(grad); sf.getparent().remove(sf)
    slide.shapes.add_picture(os.path.join(ASSETS, "logo-white.png"), E(64 + 17), E(40 + 25), E(118))
    lines = wrap(s["title"], 800, 56, W * 0.7)
    th = 62 * len(lines)
    sub = s.get("subtitle")
    y = H - 64 - th - (30 if sub else 0)
    text(slide, 64, y, W * 0.7, th, s["title"], 56, 800, "surface", lh=62, ls=-0.03)
    if sub:
        text(slide, 64, y + th + 6, W * 0.7, 24, sub, 17, 500, "E1E3E6")


def main(src, out):
    deck = json.load(open(src))
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(W), E(H)
    report = []
    for i, s in enumerate(deck["slides"], 1):
        if s.get("type") == "cover":
            build_cover(prs, s)
        else:
            build_content(prs, s, i, deck, report)
    prs.save(out)
    print("\n".join(report))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
