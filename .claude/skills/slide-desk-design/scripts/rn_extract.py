#!/usr/bin/env python3
"""Inventory a source deck so nothing gets lost in the rebuild.

    python rn_extract.py source.pptx inventory.json

Writes every text frame, table (cell by cell), chart, picture and speaker note
per slide. Pictures are flagged: a pasted table or chart needs a visual read.
The inventory is the checklist rn_check.py ticks off against the rebuilt deck.
"""
import json, sys
from pptx import Presentation
from pptx.util import Emu


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from walk(sh.shapes)
        else:
            yield sh


def main(src, out):
    prs = Presentation(src)
    slides = []
    for i, s in enumerate(prs.slides, 1):
        item = {"n": i, "hidden": s._element.get("show") == "0", "texts": [], "tables": [], "charts": [], "pictures": []}
        for sh in walk(s.shapes):
            if sh.has_text_frame and sh.text_frame.text.strip():
                size = max([r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size] or [0])
                item["texts"].append({"text": sh.text_frame.text.strip(), "pt": size, "y": round(Emu(sh.top or 0).inches, 2)})
            elif getattr(sh, "has_table", False) and sh.has_table:
                item["tables"].append([[c.text.strip() for c in r.cells] for r in sh.table.rows])
            elif getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                item["charts"].append({"type": str(ch.chart_type), "categories": [str(c) for c in ch.plots[0].categories],
                                       "series": [{"name": se.name, "values": list(se.values)} for pl in ch.plots for se in pl.series]})
            elif sh.shape_type == 13:
                item["pictures"].append({"name": sh.name, "w_in": round(Emu(sh.width).inches, 2), "h_in": round(Emu(sh.height).inches, 2),
                                         "note": "Needs a visual read if it is a pasted table or chart."})
        item["texts"].sort(key=lambda t: t["y"])
        if s.has_notes_slide:
            item["notes"] = s.notes_slide.notes_text_frame.text.strip()
        slides.append(item)
    json.dump({"source": src, "slides": slides}, open(out, "w"), indent=1, ensure_ascii=False)
    for s in slides:
        print(f"slide {s['n']}{' (hidden)' if s['hidden'] else ''}: {len(s['texts'])} text, {len(s['tables'])} tables, "
              f"{len(s['charts'])} charts, {len(s['pictures'])} pictures")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
